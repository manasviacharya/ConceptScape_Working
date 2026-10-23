import hashlib
import logging
from pathlib import Path
from typing import Optional

from pptx import Presentation
from pptx.exc import PackageNotFoundError

from ingestion.base_parser import BaseParser
from ingestion.models import IngestedResource, ContentChunk

class PPTParser(BaseParser):
    def can_handle(self, filepath: Path) -> bool:
        return filepath.suffix.lower() in [".ppt", ".pptx"]

    def parse(self, filepath: Path, resource_id: str, subject_id: str) -> Optional[IngestedResource]:
        try:
            prs = Presentation(filepath)
            chunks = []
            
            for slide_num, slide in enumerate(prs.slides, start=1):
                slide_text = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        slide_text.append(shape.text_frame.text)
                
                text = "\n".join(slide_text).strip()
                if not text:
                    logging.warning(f"Empty text found on slide {slide_num} of {filepath.name}")
                
                chunk = ContentChunk(
                    chunk_id=f"{resource_id}_slide_{slide_num}",
                    source_type="ppt",
                    location_type="slide",
                    location_number=slide_num,
                    text=text,
                    content_hash=hashlib.sha256(text.encode("utf-8")).hexdigest(),
                    extraction_method="native_text",
                )
                chunks.append(chunk)
                
            return IngestedResource(
                resource_id=resource_id,
                resource_type="ppt",
                filename=filepath.name,
                subject_id=subject_id,
                content=chunks
            )
            
        except PackageNotFoundError as e:
            logging.error(f"Error opening or reading PPT/PPTX {filepath.name}: {e}")
            return None
        except Exception as e:
            logging.error(f"Unexpected error parsing PPT/PPTX {filepath.name}: {e}")
            return None
